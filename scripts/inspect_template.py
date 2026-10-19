#!/usr/bin/env python3
import argparse,json
from pptx import Presentation
EMU=914400
def main():
 p=argparse.ArgumentParser(description="Inspect PPTX slide and shape geometry/text.");p.add_argument('--template',required=True);p.add_argument('--json');a=p.parse_args();prs=Presentation(a.template)
 out={'file':a.template,'slide_count':len(prs.slides),'size_inches':[round(prs.slide_width/EMU,2),round(prs.slide_height/EMU,2)],'slides':[]}
 for i,sl in enumerate(prs.slides,1):
  ss=[]
  for j,s in enumerate(sl.shapes):
   d={'index':j,'name':s.name,'type':str(s.shape_type),'x':round(s.left/EMU,2),'y':round(s.top/EMU,2),'w':round(s.width/EMU,2),'h':round(s.height/EMU,2)}
   if getattr(s,'has_text_frame',False):d['text']=s.text[:400]
   ss.append(d)
  out['slides'].append({'slide':i,'layout':sl.slide_layout.name,'shapes':ss})
 print(json.dumps(out,ensure_ascii=False,indent=2))
 if a.json:open(a.json,'w',encoding='utf-8').write(json.dumps(out,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
