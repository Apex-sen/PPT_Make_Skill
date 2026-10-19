#!/usr/bin/env python3
"""Append duplicated template slides and add editable text/image items from a JSON plan.
Plan: {"source_template":"template.pptx","output":"out.pptx","slides":[{"source_slide":2,"replace":{"old":"new"},"items":[{"type":"text","text":"...","x":1,"y":1,"w":4,"h":1,"font_size":18}]}]}
Coordinates are inches. Complex SmartArt, animations and OLE need manual QA.
"""
import argparse,json
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor

def clone_slide(prs,idx):
 src=prs.slides[idx];dst=prs.slides.add_slide(src.slide_layout)
 for sh in list(dst.shapes):sh.element.getparent().remove(sh.element)
 relmap={}
 for rel in src.part.rels.values():
  if rel.reltype.endswith('/slideLayout'):continue
  try:relmap[rel.rId]=dst.part.rels.add_relationship(rel.reltype,rel.target_ref if rel.is_external else rel.target_part,rel.is_external)
  except Exception:pass
 for sh in src.shapes:
  el=deepcopy(sh.element)
  for node in el.iter():
   for attr in ('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed','{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id','{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link'):
    rid=node.get(attr)
    if rid in relmap:node.set(attr,relmap[rid])
  dst.shapes._spTree.insert_element_before(el,'p:extLst')
 return dst

def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--plan',required=True);p.add_argument('--output');a=p.parse_args();plan=json.loads(Path(a.plan).read_text(encoding='utf-8'));prs=Presentation(plan['source_template'])
 for page in plan.get('slides',[]):
  n=int(page['source_slide'])
  if not 1<=n<=len(prs.slides):raise ValueError(f'source_slide must be in 1..{len(prs.slides)}')
  sl=clone_slide(prs,n-1)
  for old,new in page.get('replace',{}).items():
   for sh in sl.shapes:
    if getattr(sh,'has_text_frame',False) and old in sh.text:sh.text=sh.text.replace(old,new)
  for it in page.get('items',[]):
   x,y,w,h=[float(it[k]) for k in ('x','y','w','h')];typ=it.get('type','text')
   if typ=='image':
    if not Path(it['path']).exists():raise FileNotFoundError(it['path'])
    sl.shapes.add_picture(it['path'],Inches(x),Inches(y),width=Inches(w),height=Inches(h))
   elif typ=='text':
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.text_frame.word_wrap=True;sh.text_frame.text=it.get('text','')
    for par in sh.text_frame.paragraphs:
     for run in par.runs:
      run.font.size=Pt(float(it.get('font_size',18)));run.font.bold=bool(it.get('bold',False))
      if it.get('color'):
       c=it['color'].lstrip('#');run.font.color.rgb=RGBColor(int(c[:2],16),int(c[2:4],16),int(c[4:6],16))
   else:raise ValueError(f'unsupported item type {typ}')
 out=a.output or plan.get('output','output.pptx');Path(out).parent.mkdir(parents=True,exist_ok=True);prs.save(out);print(f'Saved {out}; slides={len(prs.slides)}')
if __name__=='__main__':main()
