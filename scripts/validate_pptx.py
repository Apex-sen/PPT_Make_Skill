#!/usr/bin/env python3
"""Check PPTX ZIP integrity, empty slides, and shapes outside slide bounds."""
import argparse,json,zipfile
from pathlib import Path
from pptx import Presentation
EMU=914400
def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--pptx',required=True);p.add_argument('--json');a=p.parse_args();path=Path(a.pptx)
 with zipfile.ZipFile(path) as z:
  bad=z.testzip()
  if bad:raise SystemExit('Corrupt ZIP member: '+bad)
 prs=Presentation(path);issues=[];slides=[]
 for i,sl in enumerate(prs.slides,1):
  shapes=[]
  for sh in sl.shapes:
   if sh.left < -1000 or sh.top < -1000 or sh.left+sh.width > prs.slide_width+1000 or sh.top+sh.height > prs.slide_height+1000:issues.append({'slide':i,'shape':sh.name,'issue':'outside_slide_bounds'})
   shapes.append({'name':sh.name,'x':round(sh.left/EMU,2),'y':round(sh.top/EMU,2),'w':round(sh.width/EMU,2),'h':round(sh.height/EMU,2)})
  if not sl.shapes:issues.append({'slide':i,'issue':'empty_slide'})
  slides.append({'slide':i,'shape_count':len(sl.shapes),'shapes':shapes})
 report={'file':str(path),'slide_count':len(prs.slides),'issues':issues,'slides':slides};print(json.dumps(report,ensure_ascii=False,indent=2))
 if a.json:Path(a.json).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
 if issues:raise SystemExit(1)
if __name__=='__main__':main()
